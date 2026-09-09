import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from megamarket.storage.report import ExcelReport, merge_excel_reports


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_report(path: Path, headers: list[str], rows: list[list]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = ExcelReport.SHEET_TITLE
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def read_rows(path: Path | io.BytesIO) -> list[tuple]:
    sheet = load_workbook(path)[ExcelReport.SHEET_TITLE]
    return list(sheet.iter_rows(values_only=True))


class MergeReportsTests(unittest.TestCase):
    def test_columns_are_matched_by_name_not_by_position(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            # Отчёт до появления колонки «Запрос», отчёт после и отчёт с
            # переставленными колонками — всё это один и тот же набор данных.
            old = make_report(
                root / "old.xlsx",
                ["Название", "Цена", "Продавец"],
                [["Дрель", "100 ₽", "Кувалда"]],
            )
            new = make_report(
                root / "new.xlsx",
                ["Название", "Цена", "Продавец", "Запрос"],
                [["iPhone", "99 000 ₽", "СОТОМАРКЕТ", "iphone 17"]],
            )
            shuffled = make_report(
                root / "shuffled.xlsx",
                ["Продавец", "Запрос", "Название"],
                [["Мвидео", "samsung s24", "Samsung"]],
            )
            target = root / "out" / "joined.xlsx"

            written = merge_excel_reports([old, new, shuffled], target)

            self.assertEqual(written, 3)
            self.assertEqual(
                read_rows(target),
                [
                    ("Название", "Цена", "Продавец", "Запрос"),
                    ("Дрель", "100 ₽", "Кувалда", None),
                    ("iPhone", "99 000 ₽", "СОТОМАРКЕТ", "iphone 17"),
                    ("Samsung", None, "Мвидео", "samsung s24"),
                ],
            )

    def test_repeated_rows_are_kept(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            headers = ["Название", "Цена"]
            row = ["Дрель", "100 ₽"]
            first = make_report(root / "first.xlsx", headers, [row])
            second = make_report(root / "second.xlsx", headers, [row])
            target = root / "joined.xlsx"

            # Повторы не ищем намеренно: объединение должно быть быстрым.
            self.assertEqual(merge_excel_reports([first, second], target), 2)
            self.assertEqual(
                read_rows(target),
                [("Название", "Цена"), ("Дрель", "100 ₽"), ("Дрель", "100 ₽")],
            )

    def test_empty_rows_and_empty_files_are_skipped(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with_gaps = make_report(
                root / "gaps.xlsx",
                ["Название", "Цена"],
                [["Дрель", "100 ₽"], [None, None], ["", ""]],
            )
            headers_only = make_report(root / "empty.xlsx", ["Название", "Цена"], [])
            target = root / "joined.xlsx"

            self.assertEqual(merge_excel_reports([with_gaps, headers_only], target), 1)
            self.assertEqual(
                read_rows(target),
                [("Название", "Цена"), ("Дрель", "100 ₽")],
            )

    def test_files_without_headers_are_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blank = make_report(root / "blank.xlsx", [], [])
            with self.assertRaisesRegex(ValueError, "заголовки"):
                merge_excel_reports([blank], root / "joined.xlsx")


class JoinEndpointTests(unittest.TestCase):
    def setUp(self):
        from megamarket.api import app as app_module

        async def no_frontend():
            return None

        patcher = patch.object(
            app_module,
            "prepare_frontend",
            new=no_frontend,
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        self.client = TestClient(app_module.app)

    @staticmethod
    def _upload(name: str, headers: list[str], rows: list[list]):
        with tempfile.TemporaryDirectory() as directory:
            path = make_report(Path(directory) / name, headers, rows)
            return ("files", (name, path.read_bytes(), XLSX))

    @staticmethod
    def _temp_directories() -> set[Path]:
        root = Path(tempfile.gettempdir())
        return set(root.glob("join-reports-*"))

    def test_uploaded_reports_come_back_as_one_file(self):
        uploads = [
            self._upload("makita.xlsx", ["Название", "Запрос"], [["Дрель", "makita"]]),
            self._upload("iphone.xlsx", ["Название", "Запрос"], [["iPhone", "iphone 17"]]),
        ]
        before = self._temp_directories()

        with self.client as client:
            response = client.post("/join", files=uploads)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["X-Files-Joined"], "2")
        self.assertEqual(response.headers["X-Rows-Joined"], "2")
        self.assertIn("all-megamarket-", response.headers["Content-Disposition"])
        self.assertEqual(
            read_rows(io.BytesIO(response.content)),
            [
                ("Название", "Запрос"),
                ("Дрель", "makita"),
                ("iPhone", "iphone 17"),
            ],
        )
        # Временная папка живёт только до конца отдачи.
        self.assertEqual(self._temp_directories(), before)

    def test_single_file_is_not_a_merge(self):
        with self.client as client:
            response = client.post(
                "/join",
                files=[self._upload("one.xlsx", ["Название"], [["Дрель"]])],
            )

        self.assertEqual(response.status_code, 422)
        self.assertIn("два файла", response.json()["detail"])

    def test_foreign_extension_is_rejected(self):
        with self.client as client:
            response = client.post(
                "/join",
                files=[
                    self._upload("good.xlsx", ["Название"], [["Дрель"]]),
                    ("files", ("отчёт.txt", b"nope", "text/plain")),
                ],
            )

        self.assertEqual(response.status_code, 422)
        self.assertIn(".xlsx", response.json()["detail"])
        self.assertFalse(self._temp_directories())


if __name__ == "__main__":
    unittest.main()
